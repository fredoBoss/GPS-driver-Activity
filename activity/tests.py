import datetime as dt
import io
import shutil
import tempfile
from pathlib import Path
from unittest import skipUnless

from django.conf import settings
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase, override_settings
from django.urls import reverse
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation

from . import excel_com, excel_log, export, log_writer, sinotrack

D1, D2 = dt.date(2026, 10, 1), dt.date(2026, 10, 2)
LOG_HEADER = ['Date', 'Driver Name', 'Plate No.', 'From (Origin)', 'Depart Time',
              'To (Destination)', 'Arrive Time', 'Park Count', 'Distance (km)', 'Remarks / DR No.']


def T(h, m):
    return dt.time(h, m)


def minutes(td):
    return None if td is None else round(td.total_seconds() / 60)


def make_workbook(path, rows, drivers=(), vehicles=(), threshold=dt.time(0, 30), header=LOG_HEADER,
                  formula_rows=0):
    """Write a workbook laid out like Driver_Activity_Log.xlsx (headers on row 4 / row 6).

    `formula_rows` prepares that many rows with the template's grey formula columns K-O,
    dropdown, highlighting and filter, like the real file's 2,000 prepared rows.
    """
    wb = Workbook()
    log = wb.active
    log.title = 'Activity Log'
    log['A1'] = 'Delivery Driver Activity Log'
    for c, title in enumerate(header, start=1):
        log.cell(row=4, column=c, value=title)
    for r, values in enumerate(rows, start=5):
        for c, value in enumerate(values, start=1):
            log.cell(row=r, column=c, value=value)
    if formula_rows:
        last = 4 + formula_rows
        for r in range(5, last + 1):
            log[f'A{r}'].number_format = 'mmm d, yyyy'
            log[f'K{r}'] = f'=IF(OR(E{r}="",G{r}=""),"",MOD(G{r}-E{r},1))'
            log[f'L{r}'] = (f'=IF(OR(G{r}="",A{r + 1}="",E{r + 1}=""),"",IF(AND(A{r + 1}=A{r},B{r + 1}=B{r},'
                            f'C{r + 1}=C{r}),IF(E{r + 1}>=G{r},E{r + 1}-G{r},"Check order"),""))')
            log[f'M{r}'] = f'=IF(ISNUMBER(L{r}),IF(L{r}>=Lists!$B$3,"LONG STAY",""),"")'
            log[f'N{r}'] = f'=IF(OR(A{r}="",B{r}=""),"",INT(A{r})&"|"&B{r})'
            log[f'O{r}'] = f'=IF(OR(E{r}="",K{r}=""),"",E{r}+K{r})'
        dv = DataValidation(type='date', operator='greaterThan', formula1='43831')
        log.add_data_validation(dv)
        dv.add(f'A5:A{last}')
        log.conditional_formatting.add(f'L5:M{last}', FormulaRule(formula=['$M5="LONG STAY"']))
        log.auto_filter.ref = f'A4:M{last}'
    lists = wb.create_sheet('Lists')
    lists['A3'] = 'Long-stay threshold (h:mm)'
    lists['B3'] = threshold
    for col, title in zip('ABDEF', ['Driver Name', 'Contact No.', 'Plate No.',
                                    'Vehicle Description', 'SinoTrack Device ID']):
        lists[f'{col}6'] = title
    for i, (name, contact) in enumerate(drivers, start=7):
        lists[f'A{i}'], lists[f'B{i}'] = name, contact
    for i, (plate, desc, device) in enumerate(vehicles, start=7):
        lists[f'D{i}'], lists[f'E{i}'], lists[f'F{i}'] = plate, desc, device
    wb.save(path)


MARK_DAY = [
    (D1, 'Mark Reyes', 'ABC 1234', 'Warehouse', T(8, 5), 'Customer A', T(8, 50), 2, 18.4, 'DR 1021'),
    (D1, 'Mark Reyes', 'ABC 1234', 'Customer A', T(9, 10), 'Customer B', T(9, 45), 1, 12.1, 'DR 1022'),
    (D1, 'Mark Reyes', 'ABC 1234', 'Customer B', T(11, 0), 'Warehouse', T(11, 50), 0, 20.6, None),
]


class WorkbookTestCase(SimpleTestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, self.tmp, ignore_errors=True)

    def build(self, rows, **kwargs):
        path = self.tmp / f'log{len(list(self.tmp.iterdir()))}.xlsx'
        make_workbook(path, rows, **kwargs)
        return path

    def load(self, rows, **kwargs):
        return excel_log.load(self.build(rows, **kwargs))


class ExcelLogTests(WorkbookTestCase):
    def test_travel_stay_and_long_stay_match_the_workbook_example(self):
        log = self.load(MARK_DAY, drivers=[('Mark Reyes', '0917 000 0000')])
        legs = log.find_driver('mark reyes').legs
        self.assertEqual([minutes(leg.travel) for leg in legs], [45, 35, 50])
        self.assertEqual([minutes(leg.stay) for leg in legs], [20, 75, None])
        self.assertEqual([leg.long_stay for leg in legs], [False, True, False])

    def test_daily_summary(self):
        day = self.load(MARK_DAY).find_driver('Mark Reyes').days[0]
        self.assertEqual(minutes(day.first_depart), 8 * 60 + 5)
        self.assertEqual(minutes(day.last_arrive), 11 * 60 + 50)
        self.assertEqual(minutes(day.time_on_road), 3 * 60 + 45)
        self.assertEqual(minutes(day.travel), 2 * 60 + 10)
        self.assertEqual(minutes(day.stay), 60 + 35)
        self.assertEqual(day.long_stays, 1)
        self.assertAlmostEqual(day.distance, 51.1)
        self.assertEqual(day.park_count, 3)

    def test_check_order_and_leg_past_midnight(self):
        log = self.load([
            (D1, 'Ana Cruz', 'XYZ 5678', 'Warehouse', T(13, 0), 'Customer C', T(13, 40)),
            (D1, 'Ana Cruz', 'XYZ 5678', 'Customer C', T(13, 20), 'Warehouse', T(14, 0)),
            (D1, 'Ana Cruz', 'XYZ 5678', 'Warehouse', T(23, 30), 'Customer D', T(0, 20)),
        ])
        legs = log.find_driver('Ana Cruz').legs
        self.assertTrue(legs[0].check_order)
        self.assertIsNone(legs[0].stay)
        self.assertEqual(minutes(legs[2].travel), 50)
        day = log.find_driver('Ana Cruz').days[0]
        self.assertEqual(minutes(day.last_arrive), 24 * 60 + 20)
        self.assertEqual(minutes(day.time_on_road), 11 * 60 + 20)

    def test_stay_needs_next_row_with_same_date_driver_and_plate(self):
        log = self.load([
            (D1, 'Mark Reyes', 'ABC 1234', 'Warehouse', T(8, 0), 'Customer A', T(8, 30)),
            (D1, 'Ana Cruz', 'XYZ 5678', 'Warehouse', T(9, 0), 'Customer C', T(9, 30)),
            (D1, 'Ana Cruz', 'DEF 9012', 'Customer C', T(10, 0), 'Warehouse', T(10, 30)),
            (D2, 'Ana Cruz', 'DEF 9012', 'Warehouse', T(11, 0), 'Customer E', T(11, 30)),
        ])
        self.assertEqual([leg.stay for d in log.drivers for leg in d.legs], [None] * 4)

    def test_drivers_from_lists_and_unlisted_names_from_the_log(self):
        log = self.load(MARK_DAY, drivers=[('Leo Tan', ''), ('Mark Reyes', '')])
        self.assertEqual([(d.name, d.in_lists, len(d.legs)) for d in log.drivers],
                         [('Leo Tan', True, 0), ('Mark Reyes', True, 3)])
        log = self.load(MARK_DAY, drivers=[('Leo Tan', '')])
        self.assertEqual([(d.name, d.in_lists) for d in log.drivers],
                         [('Leo Tan', True), ('Mark Reyes', False)])

    def test_rows_without_date_or_driver_are_reported(self):
        log = self.load([MARK_DAY[0], (None, 'Mark Reyes', 'ABC 1234', 'X', T(9, 0)), (), MARK_DAY[1]])
        self.assertEqual(log.skipped_rows, [6])
        self.assertEqual(len(log.find_driver('Mark Reyes').legs), 2)

    def test_columns_are_found_by_header(self):
        header = LOG_HEADER[:7] + ['Purpose'] + LOG_HEADER[8:]
        rows = [r[:7] + ('Delivery',) + r[8:] for r in MARK_DAY]
        log = self.load(rows, header=header)
        self.assertTrue(log.has_purpose)
        self.assertFalse(log.has_park_count)
        self.assertEqual(log.find_driver('Mark Reyes').legs[0].purpose, 'Delivery')

    def test_threshold_comes_from_lists_b3(self):
        log = self.load(MARK_DAY, threshold=dt.time(1, 30))
        self.assertEqual(minutes(log.threshold), 90)
        self.assertFalse(any(leg.long_stay for leg in log.find_driver('Mark Reyes').legs))

    def test_missing_required_column(self):
        with self.assertRaisesMessage(excel_log.ActivityLogError, '"Depart Time"'):
            self.load(MARK_DAY, header=[h if h != 'Depart Time' else 'Left' for h in LOG_HEADER])

    def test_missing_file(self):
        with self.assertRaisesMessage(excel_log.ActivityLogError, 'Cannot find'):
            excel_log.load(self.tmp / 'nope.xlsx')

    @skipUnless((settings.BASE_DIR / 'Driver_Activity_Log.xlsx').exists(), 'project workbook not present')
    def test_project_workbook_loads(self):
        excel_log.load(settings.BASE_DIR / 'Driver_Activity_Log.xlsx')


class FilterTests(SimpleTestCase):
    def test_hms(self):
        from .templatetags.activity_extras import hms
        self.assertEqual(hms(dt.timedelta(hours=1, minutes=14, seconds=5)), '01:14:05')
        self.assertEqual(hms(dt.timedelta(minutes=30)), '00:30:00')
        self.assertEqual(hms(dt.timedelta(hours=26, seconds=0.6)), '26:00:01')
        self.assertEqual(hms(None), '—')


class ViewTests(WorkbookTestCase):
    def setUp(self):
        super().setUp()
        rows = MARK_DAY + [
            (D2, 'Mark Reyes', 'ABC 1234', 'Warehouse', T(7, 0), 'Customer E', T(7, 30), 1, 5.0),
            (D1, 'Ana Cruz', 'XYZ 5678', 'Warehouse', T(13, 0), 'Customer C', T(13, 40), 1, 9.0),
        ]
        path = self.build(rows, drivers=[('Mark Reyes', '0917 000 0000'), ('Ana Cruz', ''), ('Leo Tan', '')],
                          vehicles=[('ABC 1234', 'Isuzu Elf', '9171234567')])
        self.enterContext(override_settings(ACTIVITY_LOG_PATH=path))

    def test_drivers_activity_lists_every_driver(self):
        response = self.client.get(reverse('activity:drivers_activity'))
        self.assertEqual(response.status_code, 200)
        self.assertTemplateUsed(response, 'activity/drivers_Activity.html')
        for name in ('Mark Reyes', 'Ana Cruz', 'Leo Tan', '0917 000 0000', 'ABC 1234'):
            self.assertContains(response, name)
        self.assertContains(response, reverse('activity:driver_travel_record', kwargs={'name': 'Mark Reyes'}))

    def test_drivers_activity_filters(self):
        response = self.client.get(reverse('activity:drivers_activity'), {'start': '2026-10-02', 'end': '2026-10-02'})
        mark = next(d for d in response.context['drivers'] if d.name == 'Mark Reyes')
        self.assertEqual(len(mark.legs), 1)
        self.assertContains(response, 'start=2026-10-02')  # filter carried into driver links
        response = self.client.get(reverse('activity:drivers_activity'), {'q': 'ana'})
        self.assertEqual([d.name for d in response.context['drivers']], ['Ana Cruz'])
        response = self.client.get(reverse('activity:drivers_activity'), {'start': 'soon'})
        self.assertContains(response, 'not a valid date')

    def test_driver_travel_record(self):
        url = reverse('activity:driver_travel_record', kwargs={'name': 'Mark Reyes'})
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        self.assertTemplateUsed(response, 'activity/Driver_travel_record.html')
        self.assertEqual([d.date for d in response.context['days']], [D2, D1])  # newest first
        for text in ('Customer B', '8:05 AM', '11:50 AM', '01:15:00 <span class="pill pill-danger">LONG STAY',
                     'Long stay = 00:30:00', 'Isuzu Elf', '9171234567'):
            self.assertContains(response, text)
        self.assertNotContains(response, 'Customer C')  # Ana's leg

    def test_driver_travel_record_date_filter(self):
        url = reverse('activity:driver_travel_record', kwargs={'name': 'Mark Reyes'})
        response = self.client.get(url, {'start': '2026-10-02'})
        self.assertEqual([d.date for d in response.context['days']], [D2])

    def test_unknown_driver_is_404(self):
        response = self.client.get(reverse('activity:driver_travel_record', kwargs={'name': 'Nobody'}))
        self.assertEqual(response.status_code, 404)

    def test_missing_workbook_shows_message(self):
        with override_settings(ACTIVITY_LOG_PATH=self.tmp / 'missing.xlsx'):
            response = self.client.get(reverse('activity:drivers_activity'))
        self.assertEqual(response.status_code, 503)
        self.assertContains(response, 'Cannot find the activity log', status_code=503)


TRAVEL_REPORT_HEADER = [
    'Device name', 'Start time', 'End time', 'Drive time', 'Drive mileage(km)', 'Max speed(km/h)',
    'Average speed(km/h)', '30-60 km/h', '60-90 km/h', '90-120 km/h', '> 120 km/h', 'Total',
    'Start address', 'End address', 'Start longitude(°)', 'Start latitude(°)', 'End longitude(°)',
    'End latitude(°)',
]
GONZAGA = 'T05_RJP 162 - Gonzaga'
GONZAGA_TRIPS = [
    (GONZAGA, '2026-09-01 08:00:00', '2026-09-01 08:20:00', '10.5', 'Hagonoy Road', 'Digos Diversion Road'),
    (GONZAGA, '2026-09-01 08:30:00', '2026-09-01 09:00:00', '12', 'Digos Diversion Road', 'Matanao Road'),
    (GONZAGA, '2026-09-01 10:00:00', '2026-09-01 10:15:00', '5', 'Matanao Road', 'Hagonoy Road'),
]


def sinotrack_csv(trips, blank_rows=3):
    """A Travel Report CSV the way SinoTrack exports it: BOM, tab-prefixed cells, padding, blank rows."""
    def line(cells):
        return ','.join(f'"\t{c}"' for c in cells) + ',,,,\r\n'
    text = line(TRAVEL_REPORT_HEADER)
    for device, start, end, km, origin, destination in trips:
        text += line([device, start, end, '10M 5S', km, '51', '34', '15', '', '', '', '19',
                      f' {origin} .', f'Davao del Sur {destination} .', '125.55', '7.06', '125.49', '7.02'])
    text += (',' * 21 + '\r\n') * blank_rows
    return text.encode('utf-8-sig')


class TravelReportTests(SimpleTestCase):
    def test_parse_cleans_sinotrack_cells(self):
        trips, invalid = sinotrack.parse_travel_report(sinotrack_csv(GONZAGA_TRIPS))
        self.assertEqual(invalid, [])
        self.assertEqual(len(trips), 3)
        trip = trips[0]
        self.assertEqual((trip.unit, trip.plate, trip.driver), ('T05', 'RJP 162', 'Gonzaga'))
        self.assertEqual(trip.start, dt.datetime(2026, 9, 1, 8, 0))
        self.assertEqual(trip.end, dt.datetime(2026, 9, 1, 8, 20))
        self.assertEqual((trip.start_coordinates, trip.end_coordinates), ('7.06, 125.55', '7.02, 125.49'))
        self.assertEqual(trip.start_address, 'Hagonoy Road')
        self.assertEqual(trip.end_address, 'Davao del Sur Digos Diversion Road')

    def test_device_names(self):
        self.assertEqual(sinotrack.parse_device_name(GONZAGA), ('T05', 'RJP 162', 'Gonzaga'))
        self.assertEqual(sinotrack.parse_device_name('Truck 9'), ('', 'Truck 9', ''))

    def test_rejects_other_csv(self):
        with self.assertRaisesMessage(sinotrack.ReportError, 'SinoTrack Travel Report'):
            sinotrack.parse_travel_report(b'Name,Value\r\nA,1\r\n')

    @skipUnless((settings.BASE_DIR / 'Travel Report_T05-GONZAGA.csv').exists(), 'sample export not present')
    def test_project_sample_export(self):
        data = (settings.BASE_DIR / 'Travel Report_T05-GONZAGA.csv').read_bytes()
        trips, invalid = sinotrack.parse_travel_report(data)
        self.assertEqual(len(trips), 23)
        self.assertEqual(invalid, [])
        self.assertEqual({t.device for t in trips}, {GONZAGA})


class ImportTests(WorkbookTestCase):
    def trips(self, rows=GONZAGA_TRIPS):
        return sinotrack.parse_travel_report(sinotrack_csv(rows))[0]

    def test_import_copies_each_trip_into_the_next_rows(self):
        path = self.build([], formula_rows=2)
        result = log_writer.import_trips(path, self.trips())
        self.assertEqual((result.added, result.first_row, result.last_row), (3, 5, 7))
        self.assertEqual((result.drivers, result.plates), (['Gonzaga'], ['RJP 162']))
        self.assertEqual(result.added_to_lists, ['driver Gonzaga', 'plate RJP 162'])
        self.assertTrue(result.backup.exists())

        wb = load_workbook(path)
        ws = wb['Activity Log']
        self.assertEqual([c.value for c in ws[5]][:9], [
            dt.datetime(2026, 9, 1), 'Gonzaga', 'RJP 162', 'Hagonoy Road', dt.time(8, 0),
            'Davao del Sur Digos Diversion Road', dt.time(8, 20), None, None])  # mileage is not imported
        self.assertEqual(ws['A7'].number_format, 'mmm d, yyyy')
        # Row 7 is past the 2 prepared rows: it gets translated formulas and longer ranges.
        self.assertEqual(ws['K7'].value, '=IF(OR(E7="",G7=""),"",MOD(G7-E7,1))')
        self.assertIn('E8>=G7', ws['L7'].value)
        self.assertIn('A7', ws.data_validations.dataValidation[0].sqref)
        self.assertEqual([str(cf.sqref) for cf in ws.conditional_formatting], ['L5:M7'])
        self.assertEqual(ws.auto_filter.ref, 'A4:M7')
        lists = wb['Lists']
        self.assertEqual((lists['A7'].value, lists['D7'].value, lists['E7'].value),
                         ('Gonzaga', 'RJP 162', 'Unit T05'))

        legs = excel_log.load(path).find_driver('Gonzaga').legs
        self.assertEqual([minutes(leg.stay) for leg in legs], [10, 60, None])
        self.assertEqual([leg.long_stay for leg in legs], [False, True, False])

    def test_reimport_skips_trips_already_in_the_log(self):
        path = self.build([], formula_rows=5)
        log_writer.import_trips(path, self.trips())
        result = log_writer.import_trips(path, self.trips())
        self.assertEqual((result.added, result.duplicates), (0, 3))
        self.assertEqual(len(excel_log.load(path).find_driver('Gonzaga').legs), 3)

    def test_appends_after_typed_rows_and_honours_overrides(self):
        path = self.build(MARK_DAY, formula_rows=10)
        result = log_writer.import_trips(path, self.trips(), driver='Juan Gonzaga', plate='RJP-162')
        self.assertEqual((result.first_row, result.last_row), (8, 10))
        log = excel_log.load(path)
        self.assertEqual(len(log.find_driver('Juan Gonzaga').legs), 3)
        self.assertEqual(len(log.find_driver('Mark Reyes').legs), 3)

    def test_device_name_without_driver_needs_a_driver(self):
        path = self.build([], formula_rows=5)
        trips = self.trips([('Truck 9', '2026-09-01 08:00:00', '2026-09-01 08:20:00', '1', 'A', 'B')])
        with self.assertRaisesMessage(excel_log.ActivityLogError, 'Cannot tell the driver'):
            log_writer.import_trips(path, trips)
        result = log_writer.import_trips(path, trips, driver='Pedro')
        self.assertEqual((result.drivers, result.plates), (['Pedro'], ['Truck 9']))

    def test_coordinates_are_copied_and_linked(self):
        path = self.build([], header=COORD_LOG_HEADER)
        log_writer.import_trips(path, self.trips())
        ws = load_workbook(path)['Activity Log']
        self.assertEqual([c.value for c in ws[5]][:9], [
            dt.datetime(2026, 9, 1), 'Gonzaga', 'RJP 162', 'Hagonoy Road', '7.06, 125.55', dt.time(8, 0),
            'Davao del Sur Digos Diversion Road', '7.02, 125.49', dt.time(8, 20)])
        self.assertEqual(ws['E5'].hyperlink.target, 'https://www.google.com/maps/search/?api=1&query=7.06,125.55')
        self.assertEqual(ws['H5'].hyperlink.target, 'https://www.google.com/maps/search/?api=1&query=7.02,125.49')

        with override_settings(ACTIVITY_LOG_PATH=path):
            page = self.client.get(reverse('activity:driver_travel_record', kwargs={'name': 'Gonzaga'}))
        self.assertContains(page, '<div>Hagonoy Road</div><a class="coords" '
                                  'href="https://www.google.com/maps/search/?api=1&amp;query=7.06,125.55"')

    def test_reimport_fills_blank_cells_of_trips_already_in_the_log(self):
        typed = [(SEP1, 'Gonzaga', 'RJP 162', 'Warehouse', None, T(h, m), 'Customer', None, T(h2, m2))
                 for (h, m), (h2, m2) in (((8, 0), (8, 20)), ((8, 30), (9, 0)), ((10, 0), (10, 15)))]
        typed[0] = typed[0][:4] + ('6.5, 125.1',) + typed[0][5:]  # a value someone typed stays as it is
        path = self.build(typed, header=COORD_LOG_HEADER)
        result = log_writer.import_trips(path, self.trips())
        self.assertEqual((result.added, result.duplicates, result.completed_rows), (0, 3, [5, 6, 7]))
        self.assertEqual(result.filled_columns, ['From Coordinates', 'To Coordinates'])
        ws = load_workbook(path)['Activity Log']
        self.assertEqual([(ws.cell(r, 4).value, ws.cell(r, 5).value, ws.cell(r, 8).value) for r in (5, 6, 7)], [
            ('Warehouse', '6.5, 125.1', '7.02, 125.49'),
            ('Warehouse', '7.06, 125.55', '7.02, 125.49'),
            ('Warehouse', '7.06, 125.55', '7.02, 125.49'),
        ])
        self.assertIsNone(ws['A8'].value)  # nothing appended
        again = log_writer.import_trips(path, self.trips())
        self.assertEqual((again.added, again.completed_rows), (0, []))

    @skipUnless(excel_com.pythoncom is not None, 'needs Windows with pywin32')
    def test_file_locked_by_another_program(self):
        import win32file
        path = self.build([], formula_rows=5)
        handle = win32file.CreateFile(str(path), win32file.GENERIC_READ, 0, None,  # 0 = no sharing
                                      win32file.OPEN_EXISTING, 0, None)
        try:
            with self.assertRaisesMessage(excel_log.ActivityLogError, 'locked by another program'):
                log_writer.import_trips(path, self.trips())
        finally:
            handle.Close()
        self.assertEqual(log_writer.import_trips(path, self.trips()).added, 3)

    def test_import_view(self):
        path = self.build([], formula_rows=5)
        upload = SimpleUploadedFile('Travel Report_T05-GONZAGA.csv', sinotrack_csv(GONZAGA_TRIPS), 'text/csv')
        with override_settings(ACTIVITY_LOG_PATH=path):
            response = self.client.post(reverse('activity:import_travel_report'), {'report': upload})
            self.assertRedirects(
                response,
                reverse('activity:driver_travel_record', kwargs={'name': 'Gonzaga'}),
                fetch_redirect_response=False)  # fetching would consume the one-time message
            page = self.client.get(response.url)
        self.assertContains(page, 'Copied 3 trips')
        self.assertContains(page, 'LONG STAY')

    def test_import_view_rejects_other_files(self):
        path = self.build([], formula_rows=5)
        upload = SimpleUploadedFile('notes.csv', b'a,b\r\n1,2\r\n')
        with override_settings(ACTIVITY_LOG_PATH=path):
            response = self.client.post(reverse('activity:import_travel_report'), {'report': upload}, follow=True)
        self.assertContains(response, 'not a SinoTrack Travel Report or Park Report')


COORD_LOG_HEADER = ['Date', 'Driver Name', 'Plate No.', 'From (Origin)', 'From Coordinates', 'Depart Time',
                    'To (Destination)', 'To Coordinates', 'Arrive Time', 'Park Count', 'Distance (km)']
PARK_LOG_HEADER = LOG_HEADER[:8] + ['Park Time', 'Park Address', 'Park Coordinates'] + LOG_HEADER[8:]
PARK_REPORT_HEADER = ['Device name', 'Start time', 'End time', 'Park time', 'Address', 'Longitude(°)', 'Latitude(°)']
SEP1 = dt.date(2026, 9, 1)
# Trips 08:00-08:20, 08:30-09:00 and 10:00-10:15 on Sep 1.
GONZAGA_LOG_ROWS = [
    (SEP1, 'Gonzaga', 'RJP 162', 'Hagonoy Road', T(8, 0), 'Digos Diversion Road', T(8, 20)),
    (SEP1, 'Gonzaga', 'RJP 162', 'Digos Diversion Road', T(8, 30), 'Matanao Road', T(9, 0)),
    (SEP1, 'Gonzaga', 'RJP 162', 'Matanao Road', T(10, 0), 'Hagonoy Road', T(10, 15)),
]
GONZAGA_PARKS = [
    (GONZAGA, '2026-09-01 08:20:00', '2026-09-01 08:28:00', 'Digos Diversion Road', '125.3543117', '6.7238467'),
    (GONZAGA, '2026-09-01 08:40:00', '2026-09-01 08:45:00', 'Matanao Crossing', '125.30', '6.70'),  # during trip 2
    (GONZAGA, '2026-09-01 09:00:00', '2026-09-01 09:20:00', 'Matanao Road', '125.2786433', '6.6923883'),
    (GONZAGA, '2026-09-01 09:25:00', '2026-09-01 09:58:00', 'Matanao Market', '125.28', '6.69'),
    (GONZAGA, '2026-09-01 10:15:00', '2026-09-01 18:00:00', 'Hagonoy Road', '125.34881', '6.6889083'),
    (GONZAGA, '2026-09-02 08:00:00', '2026-09-02 09:00:00', 'Davao City', '125.6', '7.1'),  # no trips that day
]


def sinotrack_park_csv(parks, blank_rows=2):
    """A Park Report CSV the way SinoTrack exports it."""
    def line(cells):
        return ','.join(f'"\t{c}"' for c in cells) + '\r\n'
    text = line(PARK_REPORT_HEADER)
    for device, start, end, address, longitude, latitude in parks:
        text += line([device, start, end, '1M 0S', f' {address} .', longitude, latitude])
    text += (',' * 6 + '\r\n') * blank_rows
    return text.encode('utf-8-sig')


class ParkReportTests(SimpleTestCase):
    def test_parse_park_report(self):
        parks, invalid = sinotrack.parse_park_report(sinotrack_park_csv(GONZAGA_PARKS))
        self.assertEqual((len(parks), invalid), (6, []))
        park = parks[0]
        self.assertEqual((park.plate, park.driver), ('RJP 162', 'Gonzaga'))
        self.assertEqual(park.address, 'Digos Diversion Road')
        self.assertEqual(park.duration, dt.timedelta(minutes=8))
        self.assertEqual(park.coordinates, '6.7238467, 125.3543117')  # latitude first, like Google Maps

    def test_detect_report(self):
        self.assertEqual(sinotrack.detect_report(sinotrack_park_csv(GONZAGA_PARKS)), 'park')
        self.assertEqual(sinotrack.detect_report(sinotrack_csv(GONZAGA_TRIPS)), 'travel')
        self.assertIsNone(sinotrack.detect_report(b'a,b\r\n1,2\r\n'))

    @skipUnless((settings.BASE_DIR / 'Park report_T05-GONZAGA.csv').exists(), 'sample export not present')
    def test_project_sample_park_report(self):
        data = (settings.BASE_DIR / 'Park report_T05-GONZAGA.csv').read_bytes()
        self.assertEqual(sinotrack.detect_report(data), 'park')
        parks, invalid = sinotrack.parse_park_report(data)
        self.assertEqual((len(parks), invalid), (41, []))


class ParkMergeTests(WorkbookTestCase):
    def parks(self):
        return sinotrack.parse_park_report(sinotrack_park_csv(GONZAGA_PARKS))[0]

    def test_stops_merge_into_the_trip_they_belong_to(self):
        path = self.build(GONZAGA_LOG_ROWS, header=PARK_LOG_HEADER)
        result = log_writer.merge_parks(path, self.parks())
        self.assertEqual((result.parks, result.matched, result.rows), (6, 5, [5, 6, 7]))
        self.assertEqual((result.unmatched, result.unmatched_dates), (1, [dt.date(2026, 9, 2)]))
        self.assertTrue(result.backup.exists())

        ws = load_workbook(path)['Activity Log']
        self.assertEqual([ws.cell(5, c).value for c in (8, 9, 10, 11)],
                         [1, dt.timedelta(minutes=8), 'Digos Diversion Road', '6.7238467, 125.3543117'])
        # Trip 2 gets the stop during the drive plus the two at its destination.
        self.assertEqual(ws['H6'].value, 3)
        self.assertEqual(ws['I6'].value, dt.timedelta(minutes=5 + 20 + 33))
        self.assertEqual(ws['J6'].value, 'Matanao Crossing\nMatanao Road\nMatanao Market')
        self.assertEqual(ws['K6'].value, '6.7, 125.3\n6.6923883, 125.2786433\n6.69, 125.28')
        self.assertEqual(ws['K6'].hyperlink.target, 'https://www.google.com/maps/search/?api=1&query=6.69,125.28')
        self.assertEqual(ws['I7'].value, dt.timedelta(hours=7, minutes=45))  # parked until the evening

        legs = excel_log.load(path).find_driver('Gonzaga').legs
        self.assertEqual([p.maps_url for p in legs[1].parks], [
            'https://www.google.com/maps/search/?api=1&query=6.7,125.3',
            'https://www.google.com/maps/search/?api=1&query=6.6923883,125.2786433',
            'https://www.google.com/maps/search/?api=1&query=6.69,125.28',
        ])
        self.assertEqual(legs[1].parks[2].address, 'Matanao Market')
        self.assertEqual(excel_log.load(path).find_driver('Gonzaga').park_time,
                         dt.timedelta(minutes=8 + 58 + 7 * 60 + 45))

    def test_stops_between_days(self):
        rows = GONZAGA_LOG_ROWS + [  # rows 5-7 on Sep 1, then Sep 2 (row 8) and Sep 5 (row 9)
            (dt.date(2026, 9, 2), 'Gonzaga', 'RJP 162', 'Hagonoy Road', T(9, 0), 'Digos', T(9, 30)),
            (dt.date(2026, 9, 5), 'Gonzaga', 'RJP 162', 'Digos', T(9, 0), 'Hagonoy Road', T(9, 30)),
        ]
        path = self.build(rows, header=PARK_LOG_HEADER)
        parks = sinotrack.parse_park_report(sinotrack_park_csv([
            (GONZAGA, '2026-09-01 10:15:00', '2026-09-02 07:00:00', 'Overnight', '125.3', '6.6'),
            (GONZAGA, '2026-09-02 07:30:00', '2026-09-02 08:00:00', 'Before the first trip', '125.3', '6.6'),
            (GONZAGA, '2026-09-03 10:00:00', '2026-09-03 11:00:00', 'Day with no trips', '125.3', '6.6'),
            (GONZAGA, '2026-09-05 07:00:00', '2026-09-05 08:00:00', 'After missing days', '125.3', '6.6'),
        ]))[0]
        result = log_writer.merge_parks(path, parks)
        # The Sep 2 morning stop comes before that day's first trip, so it joins the previous trip
        # (Sep 1, row 7). Sep 3-4 have no trips, so the Sep 3 and Sep 5 morning stops are reported.
        self.assertEqual((result.rows, result.unmatched_dates), ([7], [dt.date(2026, 9, 3), dt.date(2026, 9, 5)]))
        ws = load_workbook(path)['Activity Log']
        self.assertEqual(ws['J7'].value, 'Overnight\nBefore the first trip')

    def test_merging_again_rewrites_the_same_values(self):
        path = self.build(GONZAGA_LOG_ROWS, header=PARK_LOG_HEADER)
        log_writer.merge_parks(path, self.parks())
        log_writer.merge_parks(path, self.parks())
        ws = load_workbook(path)['Activity Log']
        self.assertEqual((ws['H6'].value, ws['I6'].value), (3, dt.timedelta(minutes=58)))

    def test_needs_the_park_columns(self):
        path = self.build(GONZAGA_LOG_ROWS)
        with self.assertRaisesMessage(excel_log.ActivityLogError, '"Park Time"'):
            log_writer.merge_parks(path, self.parks())

    def test_travel_record_links_coordinates_to_google_maps(self):
        path = self.build(GONZAGA_LOG_ROWS, header=PARK_LOG_HEADER)
        log_writer.merge_parks(path, self.parks())
        with override_settings(ACTIVITY_LOG_PATH=path):
            page = self.client.get(reverse('activity:driver_travel_record', kwargs={'name': 'Gonzaga'}))
        self.assertContains(page, '<th class="num">Park Time</th>', html=False)
        self.assertContains(page, '00:58:00')
        self.assertContains(page, 'href="https://www.google.com/maps/search/?api=1&amp;query=6.69,125.28" '
                                  'target="_blank" rel="noopener"')
        self.assertContains(page, 'Matanao Market')

    def test_import_view_merges_a_park_report(self):
        path = self.build(GONZAGA_LOG_ROWS, header=PARK_LOG_HEADER)
        upload = SimpleUploadedFile('Park report_T05-GONZAGA.csv', sinotrack_park_csv(GONZAGA_PARKS), 'text/csv')
        with override_settings(ACTIVITY_LOG_PATH=path):
            response = self.client.post(reverse('activity:import_travel_report'), {'report': upload})
            self.assertRedirects(
                response,
                reverse('activity:driver_travel_record', kwargs={'name': 'Gonzaga'}),
                fetch_redirect_response=False)
            page = self.client.get(response.url)
        self.assertContains(page, 'Merged 5 of 6 stops')
        self.assertContains(page, '1 of 6 stops in Park report_T05-GONZAGA.csv have no trip in the Activity Log '
                                  'on that day (Sep 2)')


FULL_LOG_HEADER = ['Date', 'Driver Name', 'Plate No.', 'From (Origin)', 'From Coordinates', 'Depart Time',
                   'To (Destination)', 'To Coordinates', 'Arrive Time', 'Park Count', 'Park Time', 'Park Address',
                   'Park Coordinates', 'Distance (km)', 'Remarks / DR No.']


class ExportTests(WorkbookTestCase):
    def setUp(self):
        super().setUp()
        self.path = self.build([], header=FULL_LOG_HEADER)  # laid out like the real workbook
        log_writer.import_trips(self.path, sinotrack.parse_travel_report(sinotrack_csv(GONZAGA_TRIPS))[0])
        log_writer.merge_parks(self.path, sinotrack.parse_park_report(sinotrack_park_csv(GONZAGA_PARKS))[0])
        self.enterContext(override_settings(ACTIVITY_LOG_PATH=self.path))

    def export(self, **params):
        response = self.client.get(reverse('activity:export_travel_record', kwargs={'name': 'Gonzaga'}), params)
        self.assertEqual(response.status_code, 200)
        return response, load_workbook(io.BytesIO(response.content))

    def test_export_travel_record(self):
        response, wb = self.export()
        self.assertEqual(response['Content-Type'],
                         'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        self.assertEqual(response['Content-Disposition'],
                         'attachment; filename="Travel record - Gonzaga - 2026-09-01.xlsx"')
        self.assertEqual(wb.sheetnames, ['Travel Record', 'Daily Summary'])

        ws = wb['Travel Record']
        self.assertEqual(ws['A1'].value, 'Travel Record: Gonzaga')
        self.assertEqual([c.value for c in ws[4]], [
            'Date', '#', 'Plate No.', 'From (Origin)', 'From Coordinates', 'Depart Time', 'To (Destination)',
            'To Coordinates', 'Arrive Time', 'Travel Time', 'Stay at Destination', 'Long Stay?', 'Park Count',
            'Park Time', 'Park Address', 'Park Coordinates', 'Distance (km)', 'Remarks / DR No.'])
        self.assertEqual([c.value for c in ws[5]], [
            dt.datetime(2026, 9, 1), 1, 'RJP 162', 'Hagonoy Road', '7.06, 125.55', dt.time(8, 0),
            'Davao del Sur Digos Diversion Road', '7.02, 125.49', dt.time(8, 20), dt.timedelta(minutes=20),
            dt.timedelta(minutes=10), None, 1, dt.timedelta(minutes=8), 'Digos Diversion Road',
            '6.7238467, 125.3543117', None, None])
        self.assertEqual(ws['E5'].hyperlink.target, 'https://www.google.com/maps/search/?api=1&query=7.06,125.55')
        self.assertEqual(ws['P6'].value, '6.7, 125.3\n6.6923883, 125.2786433\n6.69, 125.28')
        self.assertEqual((ws['K6'].number_format, ws['L6'].value), ('[h]:mm:ss', 'LONG STAY'))
        self.assertEqual(ws['A8'].value, 'Total')
        self.assertEqual((ws['J8'].value, ws['L8'].value), ('=SUM(J5:J7)', '=COUNTIF(L5:L7,"LONG STAY")'))

        days = wb['Daily Summary']
        self.assertEqual([c.value for c in days[5]], [
            dt.datetime(2026, 9, 1), 'RJP 162', 3, dt.time(8, 0), dt.time(10, 15),
            dt.timedelta(hours=2, minutes=15), dt.timedelta(minutes=65), dt.timedelta(minutes=70), 1, 5,
            dt.timedelta(minutes=8 + 58 + 465), None])

    def test_export_follows_the_date_filter(self):
        response, wb = self.export(start='2026-09-02')
        self.assertIsNone(wb['Travel Record']['A5'].value)
        self.assertEqual(response['Content-Disposition'],
                         'attachment; filename="Travel record - Gonzaga - from 2026-09-02.xlsx"')

    def test_travel_record_page_has_the_export_button(self):
        page = self.client.get(reverse('activity:driver_travel_record', kwargs={'name': 'Gonzaga'}),
                               {'start': '2026-09-01'})
        self.assertContains(page, 'href="/export/drivers/Gonzaga/?start=2026-09-01"')
        self.assertContains(page, 'Export to Excel')

    def test_unknown_driver_is_404(self):
        response = self.client.get(reverse('activity:export_travel_record', kwargs={'name': 'Nobody'}))
        self.assertEqual(response.status_code, 404)


class SecondUploadTests(WorkbookTestCase):
    """Uploading another driver's report must not hide the trips already in the log."""

    def upload(self, name, data):
        return self.client.post(reverse('activity:import_travel_report'),
                                {'report': SimpleUploadedFile(name, data, 'text/csv')})

    def test_earlier_trips_stay_visible_after_another_upload(self):
        capoy = 'T07_RLR 795 - CAPOY'
        capoy_trips = [(capoy, '2026-09-04 08:00:00', '2026-09-04 08:30:00', '1', 'Depot', 'Market'),
                       (capoy, '2026-09-30 09:00:00', '2026-09-30 09:40:00', '1', 'Market', 'Depot')]
        path = self.build([], formula_rows=5)
        with override_settings(ACTIVITY_LOG_PATH=path):
            self.upload('Travel Report_T05-GONZAGA.csv', sinotrack_csv(GONZAGA_TRIPS))
            response = self.upload('Travel Report_T07-CAPOY.csv', sinotrack_csv(capoy_trips))
            # The redirect no longer narrows the pages to the uploaded file's dates.
            self.assertEqual(response.url, reverse('activity:driver_travel_record', kwargs={'name': 'CAPOY'}))
            page = self.client.get(response.url)
            self.assertContains(page, 'Sep 4, 2026 – Sep 30, 2026')
            listing = self.client.get(reverse('activity:drivers_activity'))
        legs = {d.name: len(d.legs) for d in listing.context['drivers']}
        self.assertEqual(legs, {'Gonzaga': 3, 'CAPOY': 2})
        self.assertNotContains(listing, 'class="filter-notice"')

    def test_a_date_filter_says_so_and_points_to_other_dates(self):
        path = self.build([], formula_rows=5)
        log_writer.import_trips(path, sinotrack.parse_travel_report(sinotrack_csv(GONZAGA_TRIPS))[0])
        with override_settings(ACTIVITY_LOG_PATH=path):
            listing = self.client.get(reverse('activity:drivers_activity'), {'start': '2026-09-04', 'end': '2026-09-30'})
            record = self.client.get(reverse('activity:driver_travel_record', kwargs={'name': 'Gonzaga'}),
                                     {'start': '2026-09-04'})
        self.assertContains(listing, 'Showing trips from <strong>Sep 4, 2026</strong> to')
        self.assertContains(listing, '<a href="?">Show all dates</a>')
        self.assertContains(listing, '3 trips on other dates')
        self.assertContains(record, '3 trips on other dates')


CAPOY = 'T07_RLR 795 - CAPOY'
CAPOY_TRIPS = [(CAPOY, '2026-09-04 08:00:00', '2026-09-04 08:30:00', '1', 'Depot', 'Market'),
               (CAPOY, '2026-09-30 09:00:00', '2026-09-30 09:40:00', '1', 'Market', 'Depot')]


class ExportAllTests(WorkbookTestCase):
    def setUp(self):
        super().setUp()
        self.path = self.build([], header=FULL_LOG_HEADER, drivers=[('Leo Tan', '')])  # Leo has no trips
        for trips in (GONZAGA_TRIPS, CAPOY_TRIPS):
            log_writer.import_trips(self.path, sinotrack.parse_travel_report(sinotrack_csv(trips))[0])
        self.enterContext(override_settings(ACTIVITY_LOG_PATH=self.path))

    def export(self, **params):
        response = self.client.get(reverse('activity:export_all_travel_records'), params)
        self.assertEqual(response.status_code, 200)
        return response, load_workbook(io.BytesIO(response.content))

    def test_one_sheet_per_driver_named_after_the_driver(self):
        response, wb = self.export()
        self.assertEqual(response['Content-Disposition'],
                         'attachment; filename="Travel records - All drivers - 2026-09-01 to 2026-09-30.xlsx"')
        self.assertEqual(wb.sheetnames, ['Gonzaga', 'CAPOY'])  # Leo Tan has no trips, so no sheet
        gonzaga, capoy = wb['Gonzaga'], wb['CAPOY']
        self.assertEqual(gonzaga['A1'].value, 'Travel Record: Gonzaga')
        self.assertEqual([c.value for c in gonzaga[4]][:6],
                         ['Date', '#', 'Plate No.', 'From (Origin)', 'From Coordinates', 'Depart Time'])
        self.assertEqual([gonzaga[f'C{r}'].value for r in (5, 6, 7)], ['RJP 162'] * 3)
        self.assertEqual(gonzaga['A8'].value, 'Total')
        self.assertEqual(capoy['A1'].value, 'Travel Record: CAPOY')
        self.assertEqual([capoy[f'A{r}'].value for r in (5, 6, 7)],
                         [dt.datetime(2026, 9, 4), dt.datetime(2026, 9, 30), 'Total'])
        self.assertEqual((capoy['C5'].value, capoy['D5'].value), ('RLR 795', 'Depot'))
        self.assertTrue(capoy['A2'].value.startswith('Period: Sep 4, 2026 – Sep 30, 2026    Plate No.: RLR 795'))

    def test_export_follows_the_date_filter_and_the_search(self):
        response, wb = self.export(start='2026-09-02')
        self.assertEqual(wb.sheetnames, ['CAPOY'])
        self.assertEqual(response['Content-Disposition'],
                         'attachment; filename="Travel records - All drivers - 2026-09-02 to 2026-09-30.xlsx"')
        response, wb = self.export(q='gon')
        self.assertEqual(wb.sheetnames, ['Gonzaga'])
        self.assertEqual(response['Content-Disposition'],
                         'attachment; filename="Travel records - Drivers matching gon - 2026-09-01.xlsx"')

    def test_nothing_to_export(self):
        response = self.client.get(reverse('activity:export_all_travel_records'), {'start': '2027-01-01'})
        self.assertRedirects(response, reverse('activity:drivers_activity') + '?start=2027-01-01',
                             fetch_redirect_response=False)
        page = self.client.get(response.url)
        self.assertContains(page, 'Nothing to export: no driver has trips on these dates.')
        self.assertNotContains(page, 'Export all to Excel')

    def test_drivers_page_has_the_export_all_button(self):
        page = self.client.get(reverse('activity:drivers_activity'), {'start': '2026-09-01', 'q': 'gon'})
        self.assertContains(page, 'href="/export/drivers/?start=2026-09-01&amp;q=gon"')
        self.assertContains(page, 'Export all to Excel')

    def test_sheet_titles_are_valid_and_unique(self):
        taken = set()
        titles = [export.sheet_title(name, taken) for name in (
            'Gonzaga', 'GONZAGA', 'J. Cruz / Relief: [AM]', 'Maria Clara de los Santos Villanueva-Reyes',
            'Maria Clara de los Santos Villanueva-Reyes', "'Quoted'", '  ', 'History')]
        self.assertEqual(titles, [
            'Gonzaga', 'GONZAGA (2)', 'J. Cruz _ Relief_ _AM_', 'Maria Clara de los Santos Villa',
            'Maria Clara de los Santos V (2)', 'Quoted', 'Driver', 'History (2)'])
        self.assertTrue(all(len(t) <= 31 for t in titles))
