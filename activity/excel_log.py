"""Read Driver_Activity_Log.xlsx and compute the figures shown on the activity pages.

The workbook's grey formula columns are not read, because a file written by openpyxl (or
never re-saved by Excel) has no cached formula results. They are recomputed here with the
same rules as the sheet, so the website and the workbook always agree:

* Travel time = Arrive - Depart, wrapping past midnight.
* Stay        = the NEXT sheet row's Depart - this row's Arrive, only when that row has the
                same date, driver and plate. If it departs before this row arrives, the leg
                is flagged "check order" instead.
* Long stay   = stay >= the threshold in Lists!B3.
"""
import datetime as dt
import io
import re
import threading
from dataclasses import dataclass, field, replace
from functools import cached_property
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

LOG_SHEET = 'Activity Log'
LISTS_SHEET = 'Lists'
LOG_HEADER_ROW = 4
LISTS_HEADER_ROW = 6
DEFAULT_THRESHOLD = dt.timedelta(minutes=30)
DAY = dt.timedelta(days=1)

# Columns are found by their header text (case-insensitive), so renaming or moving a
# column in Excel keeps working as long as one of these names is used.
LOG_HEADERS = {
    'date': ('date',),
    'driver': ('driver name', 'driver'),
    'plate': ('plate no.', 'plate no', 'plate number', 'plate'),
    'origin': ('from (origin)', 'from', 'origin'),
    'origin_coordinates': ('from coordinates', 'start coordinates'),
    'depart': ('depart time', 'depart'),
    'destination': ('to (destination)', 'to', 'destination'),
    'destination_coordinates': ('to coordinates', 'end coordinates'),
    'arrive': ('arrive time', 'arrive'),
    'park_count': ('park count',),
    'park_time': ('park time',),
    'park_address': ('park address',),
    'park_coordinates': ('park coordinates',),
    'purpose': ('purpose',),
    'distance': ('distance (km)', 'distance'),
    'remarks': ('remarks / dr no.', 'remarks'),
}
REQUIRED_LOG_FIELDS = ('date', 'driver', 'depart', 'arrive')

LISTS_HEADERS = {
    'driver': ('driver name',),
    'contact': ('contact no.', 'contact no', 'contact'),
    'plate': ('plate no.', 'plate no', 'plate number'),
    'vehicle': ('vehicle description', 'vehicle'),
    'device_id': ('sinotrack device id', 'device id'),
}


class ActivityLogError(Exception):
    """The workbook is missing, unreadable, or not laid out like the template."""


@dataclass
class Vehicle:
    plate: str
    description: str = ''
    device_id: str = ''


@dataclass
class ParkStop:
    """One stop from the Park Address / Park Coordinates cells (one line per stop)."""
    address: str
    coordinates: str = ''

    @property
    def maps_url(self):
        return maps_url(self.coordinates)


COORDINATES = re.compile(r'(?P<lat>-?\d{1,2}(?:\.\d+)?)\s*,\s*(?P<lon>-?\d{1,3}(?:\.\d+)?)')


def maps_url(coordinates):
    """Google Maps link for a "lat, lon" cell, or '' if it holds no coordinates."""
    match = COORDINATES.search(coordinates or '')
    if not match:
        return ''
    return f'https://www.google.com/maps/search/?api=1&query={match["lat"]},{match["lon"]}'


@dataclass
class Leg:
    """One Activity Log row: the vehicle leaves `origin` and reaches `destination`.

    Times are offsets from midnight (timedelta) so legs that cross midnight can be added up.
    """
    row: int
    date: dt.date
    driver: str
    plate: str = ''
    origin: str = ''
    origin_coordinates: str = ''
    depart: dt.timedelta | None = None
    destination: str = ''
    destination_coordinates: str = ''
    arrive: dt.timedelta | None = None
    park_count: float | str | None = None
    park_time: dt.timedelta | None = None
    park_addresses: list[str] = field(default_factory=list)
    park_coordinates: list[str] = field(default_factory=list)
    purpose: str = ''
    distance: float | None = None
    remarks: str = ''
    travel: dt.timedelta | None = None
    stay: dt.timedelta | None = None
    check_order: bool = False
    long_stay: bool = False

    @property
    def arrive_after_depart(self):
        """Arrival counted from the depart day's midnight; over 24h when the leg crosses midnight."""
        if self.depart is None or self.travel is None:
            return None
        return self.depart + self.travel

    @property
    def origin_maps_url(self):
        return maps_url(self.origin_coordinates)

    @property
    def destination_maps_url(self):
        return maps_url(self.destination_coordinates)

    @property
    def parks(self):
        """Stops merged from the SinoTrack Park Report: line N of the address cell pairs with
        line N of the coordinates cell."""
        count = max(len(self.park_addresses), len(self.park_coordinates))
        return [ParkStop(self.park_addresses[i] if i < len(self.park_addresses) else '',
                         self.park_coordinates[i] if i < len(self.park_coordinates) else '')
                for i in range(count)]


@dataclass
class Day:
    """One driver's legs on one date, mirroring a row of the workbook's Daily Summary."""
    date: dt.date
    legs: list[Leg]

    @cached_property
    def plates(self):
        return _distinct(leg.plate for leg in self.legs)

    @cached_property
    def first_depart(self):
        return min((leg.depart for leg in self.legs if leg.depart is not None), default=None)

    @cached_property
    def last_arrive(self):
        return max((leg.arrive_after_depart for leg in self.legs
                    if leg.arrive_after_depart is not None), default=None)

    @cached_property
    def time_on_road(self):
        if self.first_depart is None or self.last_arrive is None:
            return None
        return self.last_arrive - self.first_depart

    @cached_property
    def travel(self):
        return _sum_durations(leg.travel for leg in self.legs)

    @cached_property
    def stay(self):
        return _sum_durations(leg.stay for leg in self.legs)

    @cached_property
    def long_stays(self):
        return sum(leg.long_stay for leg in self.legs)

    @cached_property
    def check_orders(self):
        return sum(leg.check_order for leg in self.legs)

    @cached_property
    def distance(self):
        return _sum_numbers(leg.distance for leg in self.legs)

    @cached_property
    def park_count(self):
        return _sum_numbers(leg.park_count for leg in self.legs)

    @cached_property
    def park_time(self):
        return _sum_durations(leg.park_time for leg in self.legs)


@dataclass
class Driver:
    name: str
    contact: str = ''
    in_lists: bool = True
    legs: list[Leg] = field(default_factory=list)

    def between(self, start=None, end=None):
        """A copy holding only the legs dated from `start` to `end` (inclusive; None = open)."""
        legs = [leg for leg in self.legs
                if (start is None or leg.date >= start) and (end is None or leg.date <= end)]
        return replace(self, legs=legs)

    @cached_property
    def days(self):
        """Days in date order; legs inside a day stay in sheet order."""
        by_date = {}
        for leg in self.legs:
            by_date.setdefault(leg.date, []).append(leg)
        return [Day(date, legs) for date, legs in sorted(by_date.items())]

    @cached_property
    def plates(self):
        return _distinct(leg.plate for leg in self.legs)

    @cached_property
    def last_date(self):
        return max((leg.date for leg in self.legs), default=None)

    @cached_property
    def travel(self):
        return _sum_durations(leg.travel for leg in self.legs)

    @cached_property
    def stay(self):
        return _sum_durations(leg.stay for leg in self.legs)

    @cached_property
    def long_stays(self):
        return sum(leg.long_stay for leg in self.legs)

    @cached_property
    def check_orders(self):
        return sum(leg.check_order for leg in self.legs)

    @cached_property
    def distance(self):
        return _sum_numbers(leg.distance for leg in self.legs)

    @cached_property
    def park_count(self):
        return _sum_numbers(leg.park_count for leg in self.legs)

    @cached_property
    def park_time(self):
        return _sum_durations(leg.park_time for leg in self.legs)


@dataclass
class ActivityLog:
    path: Path
    modified: dt.datetime
    threshold: dt.timedelta
    drivers: list[Driver]
    vehicles: dict[str, Vehicle]
    columns: frozenset[str]
    skipped_rows: list[int] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    @property
    def has_park_count(self):
        return 'park_count' in self.columns

    @property
    def has_trip_coordinates(self):
        return bool(self.columns & {'origin_coordinates', 'destination_coordinates'})

    @property
    def has_park_time(self):
        return 'park_time' in self.columns

    @property
    def has_park_address(self):
        return bool(self.columns & {'park_address', 'park_coordinates'})

    @property
    def has_purpose(self):
        return 'purpose' in self.columns

    def find_driver(self, name):
        key = name.strip().casefold()
        return next((d for d in self.drivers if d.name.casefold() == key), None)

    def vehicle(self, plate):
        return self.vehicles.get(plate.casefold()) or Vehicle(plate=plate)


_cache = {}
_cache_lock = threading.Lock()


def load(path):
    """Parse the workbook, reusing the last result until the file changes on disk."""
    path = Path(path)
    try:
        stat = path.stat()
    except FileNotFoundError:
        raise ActivityLogError(f'Cannot find the activity log at {path}.') from None
    signature = (stat.st_mtime_ns, stat.st_size)
    with _cache_lock:
        cached = _cache.get(path)
    if cached and cached[0] == signature:
        return cached[1]

    try:
        # Read into memory so the file is not held open (Excel may be saving it).
        data = path.read_bytes()
    except OSError as exc:
        raise ActivityLogError(
            f'Cannot read {path.name} ({exc.strerror or exc}). If Excel is saving it, '
            'wait a moment and refresh.') from exc
    log = parse(data, path=path, modified=dt.datetime.fromtimestamp(stat.st_mtime))
    with _cache_lock:
        _cache[path] = (signature, log)
    return log


def parse(data, path=Path('Driver_Activity_Log.xlsx'), modified=None):
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise ActivityLogError(f'{path.name} is not a readable Excel workbook ({exc}).') from exc
    try:
        for sheet in (LOG_SHEET, LISTS_SHEET):
            if sheet not in wb.sheetnames:
                raise ActivityLogError(f'{path.name} has no "{sheet}" sheet.')
        warnings = []
        threshold, listed_drivers, vehicles = _read_lists(wb[LISTS_SHEET], warnings)
        legs, columns, skipped = _read_legs(wb[LOG_SHEET])
    finally:
        wb.close()

    _compute_legs(legs, threshold)

    drivers = {}
    for name, contact in listed_drivers:
        drivers.setdefault(name.casefold(), Driver(name=name, contact=contact))
    for leg in legs:
        key = leg.driver.casefold()
        if key not in drivers:
            drivers[key] = Driver(name=leg.driver, in_lists=False)
        drivers[key].legs.append(leg)

    return ActivityLog(
        path=path,
        modified=modified or dt.datetime.now(),
        threshold=threshold,
        drivers=list(drivers.values()),
        vehicles=vehicles,
        columns=frozenset(columns),
        skipped_rows=skipped,
        warnings=warnings,
    )


def _compute_legs(legs, threshold):
    by_row = {leg.row: leg for leg in legs}
    for leg in legs:
        if leg.depart is not None and leg.arrive is not None:
            leg.travel = (leg.arrive - leg.depart) % DAY
        nxt = by_row.get(leg.row + 1)
        if leg.arrive is None or nxt is None or nxt.depart is None:
            continue
        if (nxt.date, nxt.driver.casefold(), nxt.plate.casefold()) != (
                leg.date, leg.driver.casefold(), leg.plate.casefold()):
            continue
        if nxt.depart >= leg.arrive:
            leg.stay = nxt.depart - leg.arrive
            leg.long_stay = leg.stay >= threshold
        else:
            leg.check_order = True


def _read_lists(ws, warnings):
    rows = list(ws.iter_rows(values_only=True))
    threshold_cell = rows[2][1] if len(rows) > 2 and len(rows[2]) > 1 else None
    threshold = _to_duration(threshold_cell)
    if threshold is None:
        threshold = DEFAULT_THRESHOLD
        warnings.append('The long-stay threshold (Lists sheet, cell B3) is blank or not a time; '
                        'using 0:30.')

    header = rows[LISTS_HEADER_ROW - 1] if len(rows) >= LISTS_HEADER_ROW else ()
    cols = match_headers(header, LISTS_HEADERS)
    drivers, vehicles = [], {}
    for values in rows[LISTS_HEADER_ROW:]:
        def get(name):
            i = cols.get(name)
            return as_text(values[i]) if i is not None and i < len(values) else ''
        if get('driver'):
            drivers.append((get('driver'), get('contact')))
        if get('plate'):
            vehicles.setdefault(get('plate').casefold(),
                                Vehicle(get('plate'), get('vehicle'), get('device_id')))
    return threshold, drivers, vehicles


def _read_legs(ws):
    rows = ws.iter_rows(min_row=LOG_HEADER_ROW, values_only=True)
    cols = match_headers(next(rows, ()), LOG_HEADERS)
    missing = [name for name in REQUIRED_LOG_FIELDS if name not in cols]
    if missing:
        expected = ', '.join(f'"{LOG_HEADERS[name][0].title()}"' for name in missing)
        raise ActivityLogError(
            f'The {LOG_SHEET} sheet is missing the column(s) {expected} in row {LOG_HEADER_ROW}.')

    legs, skipped = [], []
    for row_no, values in enumerate(rows, start=LOG_HEADER_ROW + 1):
        raw = {name: values[i] if i < len(values) else None for name, i in cols.items()}
        if all(is_blank(v) for v in raw.values()):
            continue
        date, driver = to_date(raw['date']), as_text(raw['driver'])
        if date is None or not driver:
            skipped.append(row_no)
            continue
        park_count = raw.get('park_count')
        park_number = _to_number(park_count)
        legs.append(Leg(
            row=row_no,
            date=date,
            driver=driver,
            plate=as_text(raw.get('plate')),
            origin=as_text(raw.get('origin')),
            origin_coordinates=as_text(raw.get('origin_coordinates')),
            depart=to_clock(raw['depart']),
            destination=as_text(raw.get('destination')),
            destination_coordinates=as_text(raw.get('destination_coordinates')),
            arrive=to_clock(raw['arrive']),
            park_count=park_number if park_number is not None else (as_text(park_count) or None),
            park_time=_to_duration(raw.get('park_time')),
            park_addresses=_lines(raw.get('park_address')),
            park_coordinates=_lines(raw.get('park_coordinates')),
            purpose=as_text(raw.get('purpose')),
            distance=_to_number(raw.get('distance')),
            remarks=as_text(raw.get('remarks')),
        ))
    return legs, cols, skipped


def match_headers(header, wanted):
    index = {' '.join(str(h).split()).casefold(): i
             for i, h in enumerate(header or ()) if h is not None}
    cols = {}
    for name, aliases in wanted.items():
        for alias in aliases:
            if alias in index:
                cols[name] = index[alias]
                break
    return cols


def is_blank(value):
    return value is None or (isinstance(value, str) and not value.strip())


def as_text(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return ' '.join(str(value).split())


def _to_number(value):
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return value
    if isinstance(value, str):
        try:
            return float(value.replace(',', '').strip())
        except ValueError:
            return None
    return None


def to_date(value):
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool) and value >= 1:
        return from_excel(int(value)).date()  # serial number, e.g. read from Excel as Value2
    if isinstance(value, str):
        try:
            return dt.date.fromisoformat(value.strip())
        except ValueError:
            return None
    return None


def _round_seconds(td):
    return dt.timedelta(seconds=round(td.total_seconds()))


def _to_duration(value):
    """A length of time (e.g. the long-stay threshold); Excel may hand back time or timedelta."""
    if isinstance(value, dt.timedelta):
        return _round_seconds(value)
    if isinstance(value, dt.datetime):
        value = value.time()
    if isinstance(value, dt.time):
        return _round_seconds(dt.timedelta(hours=value.hour, minutes=value.minute,
                                           seconds=value.second, microseconds=value.microsecond))
    if isinstance(value, (int, float)) and not isinstance(value, bool) and value >= 0:
        return _round_seconds(dt.timedelta(days=value))
    return None


def to_clock(value):
    """A time of day as an offset from midnight, or None if blank or not a time."""
    if isinstance(value, str):
        text = ' '.join(value.upper().split())
        for fmt in ('%H:%M', '%H:%M:%S', '%I:%M %p', '%I:%M%p', '%I:%M:%S %p', '%I %p'):
            try:
                value = dt.datetime.strptime(text, fmt).time()
                break
            except ValueError:
                continue
        else:
            return None
    duration = _to_duration(value)
    return duration % DAY if duration is not None else None


def _lines(value):
    """Non-empty lines of a multi-line cell (Alt+Enter in Excel), each with tidy spacing."""
    if is_blank(value):
        return []
    return [line for line in (as_text(part) for part in str(value).splitlines()) if line]


def _distinct(values):
    seen, result = set(), []
    for value in values:
        if value and value.casefold() not in seen:
            seen.add(value.casefold())
            result.append(value)
    return result


def _sum_durations(values):
    return sum((v for v in values if v is not None), dt.timedelta())


def _sum_numbers(values):
    numbers = [v for v in values if isinstance(v, (int, float)) and not isinstance(v, bool)]
    return sum(numbers) if numbers else None
