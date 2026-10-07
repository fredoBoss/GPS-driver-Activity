"""Copy SinoTrack reports into Driver_Activity_Log.xlsx, cell by cell, the way an encoder would.

* Travel Report (`import_trips`): each trip becomes one new Activity Log row (Date, Driver Name,
  Plate No., From, From Coordinates, Depart Time, To, To Coordinates, Arrive Time). A trip already
  in the log (same date, plate, depart and arrive time) is not added again; only its blank cells
  are filled in, so importing a file twice is harmless and new columns get completed.
* Park Report (`merge_parks`): each stop is merged into the existing trip row it belongs to,
  filling Park Count, Park Time, Park Address and Park Coordinates. Re-merging rewrites the
  same values.

A backup copy of the workbook is saved before anything is written. If the workbook is closed
it is edited with openpyxl. If Excel has it open (Windows then refuses other writers), the same
cells are typed into the open workbook through Excel and Excel saves it; see excel_com.py.
"""
import bisect
import datetime as dt
import os
import shutil
from collections import OrderedDict, defaultdict
from copy import copy
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.formatting.formatting import ConditionalFormatting
from openpyxl.formula.translate import Translator
from openpyxl.utils.datetime import to_excel
from openpyxl.worksheet.cell_range import CellRange

from . import excel_com, excel_log
from .excel_log import ActivityLogError

FIRST_ROW = excel_log.LOG_HEADER_ROW + 1
# Rows covered by the DriverList / PlateList named ranges on the Lists sheet.
LISTS_FIRST_ROW, LISTS_LAST_ROW = 7, 56
XL_UP = -4162
PARK_FIELDS = ('park_count', 'park_time', 'park_address', 'park_coordinates')
LINK_FONT_COLOR = 'FF0563C1'


@dataclass
class ImportResult:
    added: int = 0
    duplicates: int = 0
    completed_rows: list[int] = field(default_factory=list)  # existing trips whose blank cells were filled
    filled_columns: list[str] = field(default_factory=list)
    first_row: int | None = None
    last_row: int | None = None
    drivers: list[str] = field(default_factory=list)
    plates: list[str] = field(default_factory=list)
    start: dt.date | None = None
    end: dt.date | None = None
    added_to_lists: list[str] = field(default_factory=list)
    backup: Path | None = None
    warnings: list[str] = field(default_factory=list)
    via_excel: bool = False


@dataclass
class ParkMergeResult:
    parks: int = 0
    matched: int = 0
    rows: list[int] = field(default_factory=list)
    drivers: list[str] = field(default_factory=list)
    start: dt.date | None = None
    end: dt.date | None = None
    unmatched_dates: list[dt.date] = field(default_factory=list)
    unmatched: int = 0
    backup: Path | None = None
    warnings: list[str] = field(default_factory=list)
    via_excel: bool = False


@dataclass
class Snapshot:
    """What the importer needs to know about the workbook, read from the file or from Excel.

    Row tuples start at column A; dates and times may be datetimes or Excel serial numbers.
    """
    log_header: tuple
    log_rows: list[tuple]  # Activity Log rows from FIRST_ROW down
    table_last: int        # last row prepared with the grey formula columns
    lists_header: tuple
    lists_rows: list[tuple]  # Lists rows LISTS_FIRST_ROW..LISTS_LAST_ROW


@dataclass
class Plan:
    result: ImportResult | ParkMergeResult
    log_cells: dict = field(default_factory=dict)   # (row, column) -> value, 1-based like Excel
    list_cells: dict = field(default_factory=dict)
    hyperlinks: dict = field(default_factory=dict)  # (row, column) -> URL
    input_cols: list[int] = field(default_factory=list)
    table_last: int = 0
    new_rows_until: int | None = None  # last new row; rows past table_last get the formulas

    @property
    def changes(self):
        return bool(self.log_cells or self.list_cells)


def import_trips(path, trips, driver=None, plate=None, backup_dir=None):
    """Append Travel Report `trips` to the Activity Log. `driver` / `plate` override the device name."""
    return _apply(path, backup_dir, lambda snapshot: _plan_trips(snapshot, trips, driver, plate, Path(path)))


def merge_parks(path, parks, plate=None, backup_dir=None):
    """Merge Park Report `parks` into the trip rows they belong to. `plate` overrides the device name."""
    return _apply(path, backup_dir, lambda snapshot: _plan_parks(snapshot, parks, plate, Path(path)))


def _apply(path, backup_dir, planner):
    path = Path(path)
    if not path.exists():
        raise ActivityLogError(f'Cannot find the activity log at {path}.')
    backup_dir = Path(backup_dir) if backup_dir else path.parent / 'backups'
    if _is_writable(path):
        return _apply_with_openpyxl(path, planner, backup_dir)

    result = excel_com.call_with_open_workbook(
        path, lambda workbook: _apply_with_excel(workbook, path, planner, backup_dir))
    if result is excel_com.NOT_OPEN:
        raise ActivityLogError(
            f'{path.name} is locked by another program, so it cannot be updated. '
            'Close it in that program, then import again.')
    return result


# ---------------------------------------------------------------- planning (shared)

def _log_columns(snapshot, path, required):
    cols = excel_log.match_headers(snapshot.log_header, excel_log.LOG_HEADERS)
    missing = [n for n in required if n not in cols]
    if missing:
        raise ActivityLogError(f'The {excel_log.LOG_SHEET} sheet of {path.name} is missing the column(s) '
                               + ', '.join(f'"{excel_log.LOG_HEADERS[n][0].title()}"' for n in missing) + '.')
    return cols


TRIP_COORDINATE_FIELDS = ('origin_coordinates', 'destination_coordinates')


def _plan_trips(snapshot, trips, driver, plate, path):
    cols = _log_columns(snapshot, path, excel_log.REQUIRED_LOG_FIELDS + ('plate',))
    existing, last_used = _existing_trips(snapshot.log_rows, cols)

    result = ImportResult()
    plan = Plan(result, input_cols=sorted(c + 1 for c in cols.values()), table_last=snapshot.table_last)
    filled = set()
    row = last_used + 1
    for trip in trips:
        trip_driver, trip_plate = driver or trip.driver, plate or trip.plate
        if not trip_driver:
            raise ActivityLogError(
                f'Cannot tell the driver from the device name "{trip.device}". '
                'Type the driver name and import again.')
        values = {
            'date': trip.start.date(),
            'driver': trip_driver,
            'plate': trip_plate,
            'origin': trip.start_address,
            'origin_coordinates': trip.start_coordinates,
            'depart': trip.start.time(),
            'destination': trip.end_address,
            'destination_coordinates': trip.end_coordinates,
            'arrive': trip.end.time(),
        }
        values = {name: value for name, value in values.items() if name in cols and value not in (None, '')}
        key = (trip.start.date(), trip_plate.casefold(),
               excel_log.to_clock(trip.start.time()), excel_log.to_clock(trip.end.time()))

        if key in existing:
            result.duplicates += 1
            target, current = existing[key]
            if current is None:  # the same trip twice in this file
                continue
            blanks = {name: value for name, value in values.items()
                      if excel_log.is_blank(current[cols[name]] if cols[name] < len(current) else None)}
            if not blanks:
                continue
            result.completed_rows.append(target)
            filled.update(blanks)
        else:
            existing[key] = (row, None)
            target, blanks = row, values
            result.added += 1
            result.first_row = result.first_row or row
            result.last_row = row
            row += 1

        for name, value in blanks.items():
            plan.log_cells[(target, cols[name] + 1)] = value
            if name in TRIP_COORDINATE_FIELDS:
                plan.hyperlinks[(target, cols[name] + 1)] = excel_log.maps_url(value)
        result.start = min(result.start or trip.start.date(), trip.start.date())
        result.end = max(result.end or trip.start.date(), trip.start.date())
        for items, value in ((result.drivers, trip_driver), (result.plates, trip_plate)):
            if value not in items:
                items.append(value)

    result.filled_columns = [excel_log.LOG_HEADERS[name][0].title() for name in sorted(filled, key=cols.get)]
    if result.added:
        plan.new_rows_until = result.last_row
    if plan.log_cells:
        plan.list_cells = _plan_list_additions(snapshot, trips, plate, result)
    return plan


@dataclass
class _LegSpan:
    row: int
    date: dt.date
    driver: str
    depart: dt.datetime
    arrive: dt.datetime


def _plan_parks(snapshot, parks, plate, path):
    """Each stop goes to the trip row with the latest departure at or before the stop's start, so a
    stop at the destination and a stop during the drive both land on that trip, and so do overnight
    stops before the next trip. The stop must start on that trip's day (or the day it arrived, or the
    next trip's day when that is the following day); stops on days with no trips stay unmatched,
    because the Travel Report for those days has not been imported."""
    cols = _log_columns(snapshot, path, ('date', 'plate', 'depart') + PARK_FIELDS)

    def get(values, name):
        i = cols.get(name)
        return values[i] if i is not None and i < len(values) else None

    legs_by_plate = defaultdict(list)
    for row_no, values in enumerate(snapshot.log_rows, start=FIRST_ROW):
        date, depart = excel_log.to_date(get(values, 'date')), excel_log.to_clock(get(values, 'depart'))
        if date is None or depart is None:
            continue
        arrive = excel_log.to_clock(get(values, 'arrive'))
        depart_at = dt.datetime.combine(date, dt.time()) + depart
        arrive_at = depart_at + ((arrive - depart) % excel_log.DAY if arrive is not None else dt.timedelta())
        span = _LegSpan(row_no, date, excel_log.as_text(get(values, 'driver')), depart_at, arrive_at)
        legs_by_plate[excel_log.as_text(get(values, 'plate')).casefold()].append(span)
    for spans in legs_by_plate.values():
        spans.sort(key=lambda s: s.depart)
    departs = {key: [s.depart for s in spans] for key, spans in legs_by_plate.items()}

    result = ParkMergeResult(parks=len(parks))
    by_row, unmatched = defaultdict(list), []
    for park in parks:
        key = (plate or park.plate).casefold()
        i = bisect.bisect_right(departs.get(key, []), park.start) - 1
        if i < 0:
            unmatched.append(park)
            continue
        leg = legs_by_plate[key][i]
        days = {leg.date, leg.arrive.date()}
        if i + 1 < len(legs_by_plate[key]):
            next_day = legs_by_plate[key][i + 1].date
            if next_day - leg.arrive.date() <= dt.timedelta(days=1):  # no trip-less days in between
                days.add(next_day)
        if park.start.date() not in days:
            unmatched.append(park)
            continue
        by_row[leg.row].append((leg, park))
        result.matched += 1

    plan = Plan(result, input_cols=sorted(c + 1 for c in cols.values()), table_last=snapshot.table_last)
    for row, pairs in sorted(by_row.items()):
        stops = [park for _, park in pairs]
        plan.log_cells[(row, cols['park_count'] + 1)] = len(stops)
        plan.log_cells[(row, cols['park_time'] + 1)] = sum((p.duration for p in stops), dt.timedelta())
        plan.log_cells[(row, cols['park_address'] + 1)] = '\n'.join(p.address for p in stops)
        coordinates = '\n'.join(p.coordinates for p in stops)
        if coordinates.strip():
            plan.log_cells[(row, cols['park_coordinates'] + 1)] = coordinates
            # A cell holds one link: use the longest stop. The website links every stop.
            longest = max((p for p in stops if p.coordinates), key=lambda p: p.duration)
            plan.hyperlinks[(row, cols['park_coordinates'] + 1)] = excel_log.maps_url(longest.coordinates)
        leg = pairs[0][0]
        result.rows.append(row)
        if leg.driver and leg.driver not in result.drivers:
            result.drivers.append(leg.driver)
        result.start = min(result.start or leg.date, leg.date)
        result.end = max(result.end or leg.date, leg.date)
    result.unmatched = len(unmatched)
    result.unmatched_dates = sorted({p.start.date() for p in unmatched})
    return plan


def _existing_trips(rows, cols):
    """{trip key: (row, row values)} for trips already in the log, and the last row with any typed value."""
    def get(values, name):
        i = cols[name]
        return values[i] if i < len(values) else None

    trips, last_used = {}, FIRST_ROW - 1
    for row_no, values in enumerate(rows, start=FIRST_ROW):
        if all(excel_log.is_blank(get(values, name)) for name in cols):
            continue
        last_used = row_no
        key = (excel_log.to_date(get(values, 'date')),
               excel_log.as_text(get(values, 'plate')).casefold(),
               excel_log.to_clock(get(values, 'depart')),
               excel_log.to_clock(get(values, 'arrive')))
        trips.setdefault(key, (row_no, values))
    return trips, last_used


def _plan_list_additions(snapshot, trips, plate, result):
    """Add new drivers / plates to the Lists sheet (it feeds the dropdowns)."""
    cols = excel_log.match_headers(snapshot.lists_header, excel_log.LISTS_HEADERS)
    cells = {}

    def add(field_name, value, vehicle=''):
        if field_name not in cols:
            return
        i = cols[field_name]
        empty_row = None
        for row, values in enumerate(snapshot.lists_rows, start=LISTS_FIRST_ROW):
            current = cells.get((row, i + 1)) or excel_log.as_text(values[i] if i < len(values) else None)
            if current.casefold() == value.casefold():
                return
            if not current and empty_row is None:
                empty_row = row
        label = 'driver' if field_name == 'driver' else 'plate'
        if empty_row is None:
            result.warnings.append(f'The Lists sheet has no free row for {label} "{value}"; add it by hand.')
            return
        cells[(empty_row, i + 1)] = value
        if vehicle and 'vehicle' in cols:
            v = cols['vehicle']
            row_values = snapshot.lists_rows[empty_row - LISTS_FIRST_ROW]
            if excel_log.is_blank(row_values[v] if v < len(row_values) else None):
                cells[(empty_row, v + 1)] = vehicle
        result.added_to_lists.append(f'{label} {value}')

    for name in result.drivers:
        add('driver', name)
    for trip_plate in result.plates:
        unit = next((t.unit for t in trips if (plate or t.plate) == trip_plate and t.unit), '')
        add('plate', trip_plate, vehicle=f'Unit {unit}' if unit else '')
    return cells


def _backup(path, backup_dir):
    backup_dir.mkdir(parents=True, exist_ok=True)
    backup = backup_dir / f'{path.stem}_{dt.datetime.now():%Y%m%d-%H%M%S}{path.suffix}'
    shutil.copy2(path, backup)
    return backup


def _check_sheets(names, path):
    for sheet in (excel_log.LOG_SHEET, excel_log.LISTS_SHEET):
        if sheet not in names:
            raise ActivityLogError(f'{path.name} has no "{sheet}" sheet.')


# ---------------------------------------------------------------- workbook closed: openpyxl

def _is_writable(path):
    try:
        with open(path, 'r+b'):
            return True
    except PermissionError:
        return False


def _apply_with_openpyxl(path, planner, backup_dir):
    wb = load_workbook(path)
    _check_sheets(wb.sheetnames, path)
    ws, lists = wb[excel_log.LOG_SHEET], wb[excel_log.LISTS_SHEET]
    formula_cols = [c.column for c in ws[FIRST_ROW]
                    if isinstance(c.value, str) and c.value.startswith('=')]
    plan = planner(Snapshot(
        log_header=tuple(c.value for c in ws[excel_log.LOG_HEADER_ROW]),
        log_rows=list(ws.iter_rows(min_row=FIRST_ROW, values_only=True)),
        table_last=_last_formula_row(ws, formula_cols),
        lists_header=tuple(c.value for c in lists[excel_log.LISTS_HEADER_ROW]),
        lists_rows=list(lists.iter_rows(min_row=LISTS_FIRST_ROW, max_row=LISTS_LAST_ROW, values_only=True)),
    ))
    if not plan.changes:
        return plan.result

    if plan.new_rows_until and plan.new_rows_until > plan.table_last:
        for row in range(plan.table_last + 1, plan.new_rows_until + 1):
            _copy_template_row(ws, row, formula_cols)
        _extend_table(ws, plan.table_last, plan.new_rows_until)
    for (row, col), value in plan.log_cells.items():
        ws.cell(row=row, column=col, value=value)
    for (row, col), url in plan.hyperlinks.items():
        cell = ws.cell(row=row, column=col)
        cell.hyperlink = url
        font = copy(cell.font)
        font.color, font.underline = LINK_FONT_COLOR, 'single'
        cell.font = font
    for (row, col), value in plan.list_cells.items():
        lists.cell(row=row, column=col, value=value)

    plan.result.backup = _backup(path, backup_dir)
    wb.calculation.fullCalcOnLoad = True  # openpyxl drops cached results; Excel recalculates
    tmp = path.with_name(f'{path.stem}.importing{path.suffix}')
    wb.save(tmp)
    try:
        os.replace(tmp, path)
    except PermissionError:
        tmp.unlink(missing_ok=True)
        raise ActivityLogError(f'{path.name} was opened by another program during the import. '
                               'Close it there, then import again.') from None
    return plan.result


def _last_formula_row(ws, formula_cols):
    if not formula_cols:
        return ws.max_row
    col = formula_cols[0]
    for row in range(ws.max_row, FIRST_ROW - 1, -1):
        value = ws.cell(row=row, column=col).value
        if isinstance(value, str) and value.startswith('='):
            return row
    return FIRST_ROW - 1


def _copy_template_row(ws, row, formula_cols):
    """Give a row past the prepared table the first data row's formats and formulas."""
    for src in ws[FIRST_ROW]:
        dst = ws.cell(row=row, column=src.column)
        dst._style = copy(src._style)
        if src.column in formula_cols:
            dst.value = Translator(src.value, origin=src.coordinate).translate_formula(dst.coordinate)


def _extend_table(ws, old_last, new_last):
    """Stretch dropdowns, highlighting and the filter that ended at `old_last` down to `new_last`."""
    for dv in ws.data_validations.dataValidation:
        for rng in list(dv.sqref.ranges):
            if rng.max_row == old_last:
                dv.sqref.add(CellRange(min_col=rng.min_col, min_row=old_last + 1,
                                       max_col=rng.max_col, max_row=new_last))

    # ConditionalFormattingList keys its rules by range, so rebuild it with the longer ranges.
    rebuilt = OrderedDict()
    for cf, rules in ws.conditional_formatting._cf_rules.items():
        ranges = []
        for rng in cf.sqref.ranges:
            if rng.max_row == old_last:
                rng = CellRange(min_col=rng.min_col, min_row=rng.min_row,
                                max_col=rng.max_col, max_row=new_last)
            ranges.append(rng.coord)
        rebuilt[ConditionalFormatting(sqref=' '.join(ranges))] = rules
    ws.conditional_formatting._cf_rules = rebuilt

    if ws.auto_filter.ref:
        rng = CellRange(ws.auto_filter.ref)
        if rng.max_row == old_last:
            ws.auto_filter.ref = CellRange(min_col=rng.min_col, min_row=rng.min_row,
                                           max_col=rng.max_col, max_row=new_last).coord


# ---------------------------------------------------------------- workbook open in Excel

def _apply_with_excel(workbook, path, planner, backup_dir):
    if workbook.ReadOnly:
        raise ActivityLogError(f'{path.name} is open read-only in Excel. Close it there, then import again.')
    _check_sheets([sheet.Name for sheet in workbook.Worksheets], path)
    ws = workbook.Worksheets(excel_log.LOG_SHEET)
    lists = workbook.Worksheets(excel_log.LISTS_SHEET)

    used = ws.UsedRange
    last_row = max(used.Row + used.Rows.Count - 1, FIRST_ROW)
    last_col = used.Column + used.Columns.Count - 1
    formula_cols = [c for c in range(1, last_col + 1) if ws.Cells(FIRST_ROW, c).HasFormula]
    lists_used = lists.UsedRange
    lists_last_col = max(lists_used.Column + lists_used.Columns.Count - 1, 1)
    plan = planner(Snapshot(
        log_header=_grid(ws, excel_log.LOG_HEADER_ROW, 1, excel_log.LOG_HEADER_ROW, last_col)[0],
        log_rows=_grid(ws, FIRST_ROW, 1, last_row, last_col),
        table_last=(ws.Cells(ws.Rows.Count, formula_cols[0]).End(XL_UP).Row if formula_cols else last_row),
        lists_header=_grid(lists, excel_log.LISTS_HEADER_ROW, 1, excel_log.LISTS_HEADER_ROW, lists_last_col)[0],
        lists_rows=_grid(lists, LISTS_FIRST_ROW, 1, LISTS_LAST_ROW, lists_last_col),
    ))
    if not plan.changes:
        return plan.result

    plan.result.backup = _backup(path, backup_dir)
    plan.result.via_excel = True
    app = workbook.Application
    screen_updating = app.ScreenUpdating
    app.ScreenUpdating = False
    try:
        if plan.new_rows_until and plan.new_rows_until > plan.table_last:
            # Copying the last prepared row brings its formulas, formats, dropdowns and highlighting.
            first_new = plan.table_last + 1
            ws.Rows(plan.table_last).Copy(ws.Range(f'{first_new}:{plan.new_rows_until}'))
            for col in plan.input_cols:
                ws.Range(ws.Cells(first_new, col), ws.Cells(plan.new_rows_until, col)).ClearContents()
        for (row, col), value in plan.log_cells.items():
            ws.Cells(row, col).Value2 = _excel_value(value)
        for (row, col), url in plan.hyperlinks.items():
            cell = ws.Cells(row, col)
            ws.Hyperlinks.Add(Anchor=cell, Address=url, TextToDisplay=str(cell.Value2))
        for (row, col), value in plan.list_cells.items():
            lists.Cells(row, col).Value2 = _excel_value(value)
        workbook.Save()
    finally:
        app.ScreenUpdating = screen_updating
    return plan.result


def _grid(ws, row1, col1, row2, col2):
    """Cell values (Value2: dates and times as serial numbers) as a list of row tuples."""
    values = ws.Range(ws.Cells(row1, col1), ws.Cells(row2, col2)).Value2
    if not isinstance(values, tuple):
        return [(values,)]
    return [tuple(r) for r in values]


def _excel_value(value):
    if isinstance(value, dt.timedelta):
        return value.total_seconds() / 86400
    if isinstance(value, (dt.date, dt.time)):
        return to_excel(value)
    return value
