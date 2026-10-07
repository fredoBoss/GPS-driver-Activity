"""Excel export of one driver's travel record: the travel record page as a spreadsheet.

Values are real Excel dates, times and durations (not text), so the file can be sorted,
filtered and summed. Totals rows are SUM/COUNTIF formulas; Excel calculates them on open.
"""
import datetime as dt
import io
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.properties import PageSetupProperties

from .excel_log import DAY, maps_url

FONT = 'Arial'
NAVY = '1F3864'
FMT_DATE = 'mmm d, yyyy'
FMT_TIME = 'h:mm AM/PM'
FMT_NEXT_DAY = 'mmm d, h:mm AM/PM'
FMT_HM = '[h]:mm'
FMT_HMS = '[h]:mm:ss'
FMT_KM = '#,##0.0'
HEADER_ROW = 4

F_BODY = Font(name=FONT, size=10)
F_BOLD = Font(name=FONT, size=10, bold=True)
F_HEADER = Font(name=FONT, size=10, bold=True, color='FFFFFF')
F_TITLE = Font(name=FONT, size=14, bold=True, color=NAVY)
F_NOTE = Font(name=FONT, size=9, italic=True, color='595959')
F_LINK = Font(name=FONT, size=10, color='0563C1', underline='single')
FILL_HEADER = PatternFill('solid', fgColor=NAVY)
FILL_LONG = PatternFill('solid', fgColor='FDE2E1')
FILL_CHECK = PatternFill('solid', fgColor='FFF4E5')
FILL_TOTAL = PatternFill('solid', fgColor='F2F2F2')
BORDER = Border(*(Side(style='thin', color='BFBFBF'),) * 4)
TOP = Alignment(vertical='top')
TOP_WRAP = Alignment(vertical='top', wrap_text=True)


def _time(offset):
    """Offset from midnight -> time of day."""
    return None if offset is None else (dt.datetime.min + offset % DAY).time()


def _lines(values):
    return '\n'.join(values) or None


class _Column:
    def __init__(self, header, width, value, fmt=None, wrap=False, link=None, total=None):
        self.header, self.width, self.value, self.fmt = header, width, value, fmt
        self.wrap, self.link, self.total = wrap, link, total


def _leg_columns(log):
    has = log.columns.__contains__
    columns = [
        _Column('Date', 13, lambda leg, n: leg.date, FMT_DATE),
        _Column('#', 5, lambda leg, n: n),
        _Column('Plate No.', 11, lambda leg, n: leg.plate or None),
        _Column('From (Origin)', 34, lambda leg, n: leg.origin or None, wrap=True),
    ]
    if has('origin_coordinates'):
        columns.append(_Column('From Coordinates', 22, lambda leg, n: leg.origin_coordinates or None,
                               link=lambda leg: leg.origin_maps_url))
    columns += [
        _Column('Depart Time', 11, lambda leg, n: _time(leg.depart), FMT_TIME),
        _Column('To (Destination)', 34, lambda leg, n: leg.destination or None, wrap=True),
    ]
    if has('destination_coordinates'):
        columns.append(_Column('To Coordinates', 22, lambda leg, n: leg.destination_coordinates or None,
                               link=lambda leg: leg.destination_maps_url))
    columns += [
        _Column('Arrive Time', 11, lambda leg, n: _time(leg.arrive), FMT_TIME),
        _Column('Travel Time', 10, lambda leg, n: leg.travel, FMT_HM, total='SUM'),
        _Column('Stay at Destination', 12,
                lambda leg, n: 'Check order' if leg.check_order else leg.stay, FMT_HMS, total='SUM'),
        _Column('Long Stay?', 11, lambda leg, n: 'LONG STAY' if leg.long_stay else None, total='LONG STAY'),
    ]
    if has('park_count'):
        columns.append(_Column('Park Count', 9, lambda leg, n: leg.park_count, '0', total='SUM'))
    if has('park_time'):
        columns.append(_Column('Park Time', 11, lambda leg, n: leg.park_time, FMT_HMS, total='SUM'))
    if has('park_address'):
        columns.append(_Column('Park Address', 40, lambda leg, n: _lines(leg.park_addresses), wrap=True))
    if has('park_coordinates'):
        columns.append(_Column('Park Coordinates', 22, lambda leg, n: _lines(leg.park_coordinates), wrap=True,
                               link=lambda leg: next((p.maps_url for p in leg.parks if p.maps_url), '')))
    if has('purpose'):
        columns.append(_Column('Purpose', 16, lambda leg, n: leg.purpose or None))
    columns += [
        _Column('Distance (km)', 11, lambda leg, n: leg.distance, FMT_KM, total='SUM'),
        _Column('Remarks / DR No.', 26, lambda leg, n: leg.remarks or None, wrap=True),
    ]
    return columns


def _day_columns(log):
    has = log.columns.__contains__

    def last_arrive(day):
        if day.last_arrive is None:
            return None
        if day.last_arrive >= DAY:  # after midnight: show the date too
            return dt.datetime.combine(day.date, dt.time()) + day.last_arrive
        return _time(day.last_arrive)

    columns = [
        _Column('Date', 13, lambda day, n: day.date, FMT_DATE),
        _Column('Plate No.', 14, lambda day, n: ', '.join(day.plates) or None),
        _Column('Trip Legs', 9, lambda day, n: len(day.legs), '0', total='SUM'),
        _Column('First Depart', 12, lambda day, n: _time(day.first_depart), FMT_TIME),
        _Column('Last Arrive', 16, lambda day, n: last_arrive(day), FMT_TIME),
        _Column('Time on Road', 11, lambda day, n: day.time_on_road, FMT_HM, total='SUM'),
        _Column('Travel Time', 11, lambda day, n: day.travel, FMT_HM, total='SUM'),
        _Column('Stay Time', 11, lambda day, n: day.stay, FMT_HMS, total='SUM'),
        _Column('Long Stays', 10, lambda day, n: day.long_stays, '0', total='SUM'),
    ]
    if has('park_count'):
        columns.append(_Column('Park Count', 10, lambda day, n: day.park_count, '0', total='SUM'))
    if has('park_time'):
        columns.append(_Column('Park Time', 11, lambda day, n: day.park_time, FMT_HMS, total='SUM'))
    columns.append(_Column('Distance (km)', 12, lambda day, n: day.distance, FMT_KM, total='SUM'))
    return columns


def travel_record_workbook(log, driver, start=None, end=None):
    """The driver's legs (already limited to start..end) as .xlsx bytes."""
    wb = Workbook()
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    days = driver.days
    period = _period(days, start, end)
    plates = ', '.join(driver.plates) or '—'
    notes = [f'Period: {period}    Plate No.: {plates}' + (f'    Contact No.: {driver.contact}' if driver.contact else ''),
             f'From {log.path.name}, exported {dt.datetime.now():%b %d, %Y %I:%M %p}. '
             f'Long stay = a stay of {_hms(log.threshold)} or more (Lists sheet, cell B3).']

    legs_ws = wb.active
    legs_ws.title = 'Travel Record'
    rows = [(leg, n, leg) for day in days for n, leg in enumerate(day.legs, start=1)]
    _write_sheet(legs_ws, f'Travel Record: {driver.name}', notes, _leg_columns(log), rows, freeze='D5')

    days_ws = wb.create_sheet('Daily Summary')
    _write_sheet(days_ws, f'Daily Summary: {driver.name}', notes, _day_columns(log),
                 [(day, n, None) for n, day in enumerate(days, start=1)], freeze='B5')

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _write_sheet(ws, title, notes, columns, rows, freeze):
    ws['A1'] = title
    ws['A1'].font = F_TITLE
    for i, note in enumerate(notes, start=2):
        ws.cell(row=i, column=1, value=note).font = F_NOTE

    for c, column in enumerate(columns, start=1):
        cell = ws.cell(row=HEADER_ROW, column=c, value=column.header)
        cell.font, cell.fill, cell.border = F_HEADER, FILL_HEADER, BORDER
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = column.width
    ws.row_dimensions[HEADER_ROW].height = 30

    first = HEADER_ROW + 1
    for r, (item, n, leg) in enumerate(rows, start=first):
        fill = None
        if leg is not None and leg.long_stay:
            fill = FILL_LONG
        elif leg is not None and leg.check_order:
            fill = FILL_CHECK
        for c, column in enumerate(columns, start=1):
            cell = ws.cell(row=r, column=c, value=column.value(item, n))
            cell.border = BORDER
            cell.alignment = TOP_WRAP if column.wrap else TOP
            cell.font = F_BODY
            if column.fmt:
                cell.number_format = column.fmt
                if isinstance(cell.value, dt.datetime) and column.fmt == FMT_TIME:
                    cell.number_format = FMT_NEXT_DAY
            if column.link and cell.value is not None:
                url = column.link(item)
                if url:
                    cell.hyperlink = url
                    cell.font = F_LINK
            if fill:
                cell.fill = fill

    last = first + len(rows) - 1
    if rows:
        total_row = last + 1
        ws.cell(row=total_row, column=1, value='Total').font = F_BOLD
        for c, column in enumerate(columns, start=1):
            cell = ws.cell(row=total_row, column=c)
            cell.fill, cell.border, cell.alignment = FILL_TOTAL, BORDER, TOP
            if not column.total:
                continue
            rng = f'{get_column_letter(c)}{first}:{get_column_letter(c)}{last}'
            if column.total == 'SUM':
                cell.value = f'=SUM({rng})'
                cell.number_format = column.fmt or 'General'
            else:
                cell.value = f'=COUNTIF({rng},"{column.total}")'
            cell.font = F_BOLD
        ws.auto_filter.ref = f'A{HEADER_ROW}:{get_column_letter(len(columns))}{last}'

    ws.freeze_panes = freeze
    ws.print_title_rows = f'{HEADER_ROW}:{HEADER_ROW}'
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)


def _period(days, start, end):
    first = start or (days[0].date if days else None)
    last = end or (days[-1].date if days else None)
    if not first and not last:
        return 'all dates'
    fmt = lambda d: f'{d:%b} {d.day}, {d:%Y}' if d else '…'
    return fmt(first) if first == last else f'{fmt(first)} – {fmt(last)}'


def _hms(td):
    minutes, seconds = divmod(round(td.total_seconds()), 60)
    hours, minutes = divmod(minutes, 60)
    return f'{hours:02d}:{minutes:02d}:{seconds:02d}'


def export_filename(driver_name, start=None, end=None, days=()):
    """e.g. "Travel record - Gonzaga - 2026-09-01 to 2026-09-03.xlsx"."""
    safe_name = re.sub(r'[^\w .-]+', '_', driver_name).strip() or 'driver'
    first = start or (days[0].date if days else None)
    last = end or (days[-1].date if days else None)
    if first and last:
        period = first.isoformat() if first == last else f'{first.isoformat()} to {last.isoformat()}'
    elif first:
        period = f'from {first.isoformat()}'
    elif last:
        period = f'until {last.isoformat()}'
    else:
        period = 'all dates'
    return f'Travel record - {safe_name} - {period}.xlsx'
